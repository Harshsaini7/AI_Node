"""
Spreadsheet Manager
====================
Handles all Excel spreadsheet operations:
- Loading the workbook
- Finding PENDING customers
- Updating customer status with color coding
- Saving the workbook

Phase 1.4 - Update with Color Coding
"""

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font


# Column positions (1-indexed)
COL_NAME = 1
COL_PHONE = 2
COL_STATUS = 3

# Header row is row 1, data starts from row 2
DATA_START_ROW = 2

# Color fills for status
FILL_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_NONE = PatternFill(fill_type=None)  # No fill for PENDING

FONT_GREEN = Font(name="Calibri", size=11, color="006100")
FONT_RED = Font(name="Calibri", size=11, color="9C0006")
FONT_DEFAULT = Font(name="Calibri", size=11)

# Valid statuses
VALID_STATUSES = ["INTERESTED", "NOT_INTERESTED", "PENDING"]


def load_spreadsheet(filepath):
    """
    Load an Excel workbook from the given filepath.

    Args:
        filepath (str): Path to the .xlsx file

    Returns:
        Workbook: The loaded openpyxl Workbook object

    Raises:
        FileNotFoundError: If the file doesn't exist
        Exception: If the file is not a valid Excel file
    """
    wb = load_workbook(filepath)
    return wb


def validate_spreadsheet(workbook):
    """
    Check that the spreadsheet has the required columns: Name, Phone Number, Status.

    Args:
        workbook: The openpyxl Workbook object

    Returns:
        bool: True if valid

    Raises:
        ValueError: If required columns are missing
    """
    ws = workbook.active
    expected_headers = ["Name", "Phone Number", "Status"]

    for col_num, expected in enumerate(expected_headers, 1):
        actual = ws.cell(row=1, column=col_num).value
        if actual != expected:
            raise ValueError(
                f"Column {col_num} should be '{expected}' but found '{actual}'. "
                f"Please check your spreadsheet format."
            )

    return True


def get_pending_customers(workbook):
    """
    Find all customers whose status is PENDING.

    Args:
        workbook: The openpyxl Workbook object

    Returns:
        list: List of dicts with keys: name, phone, row_number
              Example: [{"name": "Rahul Sharma", "phone": "9000000001", "row_number": 2}]
    """
    ws = workbook.active
    pending_customers = []

    # Iterate through all data rows (starting from row 2)
    for row_num in range(DATA_START_ROW, ws.max_row + 1):
        name = ws.cell(row=row_num, column=COL_NAME).value
        phone = ws.cell(row=row_num, column=COL_PHONE).value
        status = ws.cell(row=row_num, column=COL_STATUS).value

        # Skip rows with missing data
        if not name or not phone:
            continue

        # Check if status is PENDING
        if status and status.strip().upper() == "PENDING":
            pending_customers.append({
                "name": name.strip(),
                "phone": str(phone).strip(),
                "row_number": row_num,
            })

    return pending_customers


def update_customer_status(workbook, row_number, new_status):
    """
    Update a customer's status and apply color coding to the entire row.

    Colors:
        INTERESTED     -> Green row
        NOT_INTERESTED -> Red row
        PENDING        -> No fill (neutral)

    Args:
        workbook: The openpyxl Workbook object
        row_number (int): The row number to update (1-indexed)
        new_status (str): One of 'INTERESTED', 'NOT_INTERESTED', 'PENDING'

    Raises:
        ValueError: If the status is not valid
    """
    # Validate the status
    new_status = new_status.strip().upper()
    if new_status not in VALID_STATUSES:
        raise ValueError(
            f"Invalid status '{new_status}'. Must be one of: {VALID_STATUSES}"
        )

    ws = workbook.active

    # Update the status cell
    ws.cell(row=row_number, column=COL_STATUS, value=new_status)

    # Choose the color based on status
    if new_status == "INTERESTED":
        fill = FILL_GREEN
        font = FONT_GREEN
    elif new_status == "NOT_INTERESTED":
        fill = FILL_RED
        font = FONT_RED
    else:  # PENDING
        fill = FILL_NONE
        font = FONT_DEFAULT

    # Apply color to ALL columns in the row (Name, Phone, Status)
    for col in range(COL_NAME, COL_STATUS + 1):
        ws.cell(row=row_number, column=col).fill = fill
        ws.cell(row=row_number, column=col).font = font


def save_spreadsheet(workbook, filepath):
    """
    Save the workbook back to the file.

    Args:
        workbook: The openpyxl Workbook object
        filepath (str): Path to save the .xlsx file
    """
    workbook.save(filepath)


# ---- Run this file directly to test ----
if __name__ == "__main__":
    FILEPATH = "loan_leads.xlsx"

    print("=" * 50)
    print("  SPREADSHEET MANAGER - Test")
    print("=" * 50)
    print()

    # Test 1: Load spreadsheet
    try:
        wb = load_spreadsheet(FILEPATH)
        print("[OK] Spreadsheet loaded successfully")
    except FileNotFoundError:
        print("[ERROR] loan_leads.xlsx not found!")
        print("Run 'python create_sample_data.py' first.")
        exit(1)

    # Test 2: Validate columns
    try:
        validate_spreadsheet(wb)
        print("[OK] Spreadsheet format is valid")
    except ValueError as e:
        print(f"[ERROR] {e}")
        exit(1)

    # Test 3: Find pending customers
    pending = get_pending_customers(wb)
    print(f"[OK] Found {len(pending)} PENDING customers")

    if len(pending) >= 2:
        # Test 4: Update first customer to INTERESTED (GREEN)
        first = pending[0]
        update_customer_status(wb, first["row_number"], "INTERESTED")
        print(f"[OK] {first['name']} -> INTERESTED (GREEN)")

        # Test 5: Update second customer to NOT_INTERESTED (RED)
        second = pending[1]
        update_customer_status(wb, second["row_number"], "NOT_INTERESTED")
        print(f"[OK] {second['name']} -> NOT_INTERESTED (RED)")

        # Test 6: Save
        save_spreadsheet(wb, FILEPATH)
        print(f"[OK] Spreadsheet saved to {FILEPATH}")

        # Verify: Reload and check remaining PENDING count
        wb2 = load_workbook(FILEPATH)
        remaining = get_pending_customers(wb2)
        print(f"[OK] Remaining PENDING customers: {len(remaining)}")
    else:
        print("[SKIP] Need at least 2 PENDING customers to test updates")

    print()
    print("Phase 1.4 complete.")
    print("Open loan_leads.xlsx to see the GREEN and RED rows!")
