"""
Create Sample Spreadsheet
=========================
Generates a sample loan_leads.xlsx file with fake customer data.
All customers start with PENDING status.

Run this script to create fresh test data:
    python create_sample_data.py

Phase 1.2
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def create_sample_spreadsheet(filepath="loan_leads.xlsx"):
    """Create a sample loan leads spreadsheet with fake customer data."""

    # Create a new workbook and get the active sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Loan Leads"

    # ---- Define styles ----

    # Header style
    header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    # Border style
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # ---- Write header row ----
    headers = ["Name", "Phone Number", "Status"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # ---- Sample fake customer data ----
    # All customers start as PENDING
    customers = [
        ("Rahul Sharma", "9000000001", "PENDING"),
        ("Amit Kumar", "9000000002", "PENDING"),
        ("Priya Singh", "9000000003", "PENDING"),
        ("Neha Verma", "9000000004", "PENDING"),
        ("Vikram Patel", "9000000005", "PENDING"),
        ("Sneha Gupta", "9000000006", "PENDING"),
        ("Arjun Reddy", "9000000007", "PENDING"),
        ("Kavita Nair", "9000000008", "PENDING"),
    ]

    # ---- Write customer data ----
    data_font = Font(name="Calibri", size=11)
    data_alignment = Alignment(horizontal="center", vertical="center")

    for row_num, (name, phone, status) in enumerate(customers, 2):
        # Name column
        name_cell = ws.cell(row=row_num, column=1, value=name)
        name_cell.font = data_font
        name_cell.alignment = Alignment(horizontal="left", vertical="center")
        name_cell.border = thin_border

        # Phone Number column
        phone_cell = ws.cell(row=row_num, column=2, value=phone)
        phone_cell.font = data_font
        phone_cell.alignment = data_alignment
        phone_cell.border = thin_border

        # Status column
        status_cell = ws.cell(row=row_num, column=3, value=status)
        status_cell.font = data_font
        status_cell.alignment = data_alignment
        status_cell.border = thin_border

    # ---- Set column widths for readability ----
    ws.column_dimensions["A"].width = 22  # Name
    ws.column_dimensions["B"].width = 18  # Phone Number
    ws.column_dimensions["C"].width = 18  # Status

    # ---- Save the file ----
    wb.save(filepath)
    print(f"Spreadsheet created: {filepath}")
    print(f"Total customers: {len(customers)}")
    print(f"All customers set to: PENDING")


if __name__ == "__main__":
    create_sample_spreadsheet()
